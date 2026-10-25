"""Generate versioned Word/PDF release dossier from repository facts.
Only used in CI; generated examples are explicitly illustrative.
Author: Abdelhalim Talbi.
"""
from pathlib import Path
from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from xml.sax.saxutils import escape

root = Path(__file__).resolve().parents[1]
version = (root / "VERSION.txt").read_text(encoding="utf-8").strip()
dest = root / "05-Documentation"
dest.mkdir(exist_ok=True)
sections = [
    ("Scope", "OpenRemote Studio V" + version + " — Windows candidate. Author: Abdelhalim Talbi."),
    ("Primary workflows", "Connections, Sessions, Diagnostics and Transfers are the four main views. Ctrl+1 to Ctrl+4 navigate them."),
    ("SFTP history", "The Transfers workspace shows local, read-only transfer journal entries. Revisiting the tab refreshes the journal. Corrupt journal snapshots clear the table and show a neutral error."),
    ("Session safety", "Reconnect is rejected if local session observations are malformed or a sibling local process is still active. A local process does not prove remote connectivity."),
    ("Quality gates", "GitHub Windows CI checks Python unit tests, PowerShell syntax, PE/ZIP structure, isolated portable startup, runtime DLLs and Qt platform plugin."),
    ("Known limits", "Real-world multi-host SSH/SFTP/RDP/VNC acceptance tests, signed binaries and real UI screenshots are not evidenced by this document. No fictional screenshots are included."),
    ("Next versions", "V31.6: interrupted SFTP transfer integrity. V31.7: GUI interaction and accessibility. V31.8: structured security audit. V31.9: Windows end-to-end qualification."),
]
# Version-specific addendum, sourced from the tracked R96 design note.
# This is architecture documentation, NOT an assertion that live WinRM passed.
if (dest / "V31.8-R96-WINRM-SERVICES-QT.md").is_file():
    sections.extend([
        ("R96 - Windows services read-only", "Outils > Services Windows WinRM: a separate Qt tab displays name, display label, state and startup mode. Up to 128 entries are shown. A text filter works locally and never performs a new WinRM POST."),
        ("R96 - consent and lifecycle", "Offline by default. HTTPS certificate validation and explicit consent are mandatory. A transient masked password or injected runtime provider is used; the UI does not persist or export secrets. A generation gate rejects stale results after consent or endpoint changes, and prevents concurrent POST while a previous call drains."),
        ("R96 - quality cases", "Eight Qt tests with synthetic WinRM clients cover offline safety, local filtering, truncation, stale cancellation, redacted exceptions, malformed data, empty display names and unsupported operations. Seven backend R95 tests cover the fixed CIM query and bounded normalization. CI results must be checked against the exact commit."),
        ("R96 - unverified boundaries", "No real Windows Server WinRM HTTPS acceptance test, physical Windows operator walkthrough, vault reference integration, R96 panel screenshot, or pywinrm runtime bundling is asserted. The OpenSSH integration suite does not qualify WinRM. No simulated screenshot is presented as real."),
        ("R96 - architecture", "Qt WinRMServicesPanel -> single QRunnable with fixed operation 'services' -> WinRMServicesReadOnly -> pywinrm optional TLS-validated transport -> fixed Get-CimInstance Win32_Service -> validated JSON <=128 rows -> local QTreeWidget filter. No arbitrary PowerShell or remote writes."),
    ])
# R99 release addendum: source facts and explicit unverified boundaries.
if (dest / "V31.8-R99-WINRM-WINCRED-REFERENCE.md").is_file():
    sections.extend([
        ("R99 - WinCred reference", "The optional wincred:<target> field in both Windows inventory and services panels resolves an existing Windows Credential Manager Generic Credential inside the worker after explicit operator consent. No credential creation, export or persistence is performed by the WinRM panel."),
        ("R99 - fail closed", "The target uses a bounded ASCII allowlist. Invalid references are rejected before network access. Provider exceptions become fixed error codes; no raw error or secret reaches the UI. Editing the endpoint, account or reference revokes consent and cancels pending work."),
        ("R99 - verification", "Six pure contract tests and five Qt synthetic transport tests are included. The Windows CI must run them at the exact source SHA. Local POSIX Qt tests may be skipped if PySide6 is absent."),
        ("R99 - unverified", "No real Windows Server WinRM HTTPS acceptance test, actual Windows Credential Manager integration test on a workstation, physical Windows operator test or screenshot of the R99 reference flow is asserted. pywinrm remains optional; Windows installer inclusion is not established by these tests."),
    ])
# R100: distinguish real local WinCred QA from unverified remote WinRM.
if (dest / "V31.8-R100-WINCRED-NATIVE-BOUNDARIES.md").is_file():
    sections.extend([
        ("R100 - WinCred memory boundary", "Generic Credential blobs are rejected if the native type is wrong, the byte length is zero, odd or greater than 2560, or decoded UTF-16 contains NUL. Native memory is freed on success and rejected blobs."),
        ("R100 - real Windows CredMan test", "The Windows CI test writes a random session-only Generic Credential to the isolated runner account, reads it through the production wincred: resolver and deletes it in finally. No remote server is contacted. Results must be checked on the exact commit."),
        ("R100 - limits", "Real Windows Credential Manager testing does NOT qualify WinRM HTTPS against Windows Server, TLS trust chains, NTLM/Kerberos in production, physical Windows UI operation or pywinrm bundling. No screenshots are invented."),
    ])
# R105: local enterprise CA trust is not remote WinRM acceptance.
if (dest / "V31.8-R105-WINRM-ENTERPRISE-CA.md").is_file():
    sections.extend([
        ("R105 - enterprise CA WinRM HTTPS", "Windows inventory and services accept an optional local PEM CA trust file. The default system trust remains unchanged. Explicit CA trust is passed to pywinrm with server certificate validation enabled."),
        ("R105 - fail closed and consent", "Absolute local regular PEM only, no symlinks or UNC, maximum 1 MiB, parsed before credential access. Changing the CA revokes Qt consent. Errors expose no path or secret."),
        ("R105 - qualification limits", "Eight offline backend tests use a real public CA test certificate; three Qt tests and an actual offline UI capture run in Windows CI. No real Windows Server WinRM HTTPS, NTLM/Kerberos or operator acceptance is claimed."),
    ])
# R104: WinRM remote display fields are untrusted, including injected adapters.
if (dest / "V31.8-R104-WINRM-DISPLAY-TRUST.md").is_file():
    sections.extend([
        ("R104 - Unicode display trust", "The read-only WinRM backend rejects invisible Unicode format characters, bidirectional controls, surrogates, line separators and Unicode noncharacters in remote CIM display fields. International visible text remains supported."),
        ("R104 - Qt validation boundary", "The Windows OS/disks and services Qt panels revalidate untrusted normalized results, including custom injected adapters, before populating the UI. Invalid disks, impossible free-space counts and spoofed service names are rejected with neutral errors."),
        ("R104 - evidence and limits", "Five backend tests passed locally; three Qt tests require Windows PySide6 CI and must be checked on the exact commit. No real Windows Server WinRM HTTPS session, CA/NTLM authentication or physical Windows operator walkthrough has been qualified. No R104 screenshot is claimed."),
        ("R104 - release manifest integrity", "The first R104 CI archive exposed a stale nested source manifest despite a valid outer manifest. The cumulative builder now regenerates the source manifest from exactly archived bytes and the verifier checks both manifest file sets and SHA256 digests. Four pure tests cover this release gate; only the final CI artifact can be declared qualified."),
    ])
# R103: external WinRM exception boundary and Qt signal redaction.
if (dest / "V31.8-R103-WINRM-EXCEPTION-BOUNDARIES.md").is_file():
    sections.extend([
        ("R103 - untrusted WinRM errors", "External session factories, run_ps calls and response properties may throw exceptions containing secrets. Both read-only backends now normalize all external exceptions to fixed error codes and suppress chained provider context."),
        ("R103 - Qt signal privacy", "The Qt worker emits only allowlisted error codes. Unexpected WinRMError strings from injected clients cannot enter the signal payload. Explicit consent, verified HTTPS, offline default and fixed CIM read-only commands are unchanged."),
        ("R103 - qualification boundary", "Hostile in-memory transports are covered by backend tests; Qt signal tests require PySide6 and Windows CI. Real Windows Server WinRM HTTPS, CA/NTLM/Kerberos and physical operator UI remain unverified. No R103 screenshots are claimed."),
    ])
# R102: reject ambiguous remote WinRM responses without leaking provider errors.
if (dest / "V31.8-R102-WINRM-STRICT-RESPONSE.md").is_file():
    sections.extend([
        ("R102 - strict WinRM reply decoding", "OS/disk and Windows service inventory now share a bounded UTF-8 JSON decoder: 1-65536 bytes, unique object keys, finite JSON values and integer status code. Invalid or deeply nested payloads are rejected."),
        ("R102 - redaction and offline guarantee", "Unexpected transport response property errors are replaced with fixed error codes without exposing provider exceptions or remote stderr. No network is attempted without explicit operator consent; fixed CIM queries and TLS validation are unchanged."),
        ("R102 - qualification boundary", "Six new synthetic-transport tests plus existing R93 and R95 regression tests run locally. A real Windows Server WinRM HTTPS session, CA validation, NTLM/Kerberos in situ, physical Windows operator test and R102 screenshots remain unverified."),
    ])

# R110: Zabbix network consent is bound to the exact reviewed origin.
if (dest / "V31.8-R110-ZABBIX-CONSENT-BINDING.md").is_file():
    sections.extend([
        ("R110 - Zabbix consent binding", "Editing the HTTPS origin automatically unchecks online consent, clears active alert rows and invalidates pending work. A new origin requires a fresh explicit approval before the token provider can be read or another request can be sent."),
        ("R110 - Qt regression and screenshot", "Two Qt regression tests cover re-consent after an origin change and rejection of a late response from the old origin. Windows CI captures the real Qt consent-revoked state offline with .invalid example origins; no real token or network request is used."),
        ("R110 - limitations", "A request already inside the HTTPS transport cannot be guaranteed interrupted mid-read; its result is discarded and the request remains timeout-bounded. A real Zabbix instance, least-privilege API role, real token and production behavior remain unverified. Zabbix stays PARTIAL."),
    ])
# R111: enforce a monotonic deadline while reading Zabbix response bodies.
if (dest / "V31.8-R111-ZABBIX-RESPONSE-DEADLINE.md").is_file():
    sections.extend([
        ("R111 - bounded Zabbix body read", "Zabbix response bodies are streamed in bounded read1 chunks. A monotonic deadline and the remaining socket timeout prevent a slow-drip body from extending a read indefinitely; cancellation is rechecked between chunks and errors remain neutral."),
        ("R111 - deadline scope", "The body deadline begins after urllib has completed connection/TLS and received response headers. The earlier connection/TLS/header phase still uses urllib's socket inactivity timeout; no end-to-end absolute deadline is claimed."),
        ("R111 - qualification boundary", "Synthetic clock/socket tests plus a real local socketpair slow-drip scenario cover deadline expiry, cancellation, response-size bounds and error redaction. A real Zabbix server, real token and effective least-privilege API role remain unverified. Zabbix stays PARTIAL and no R111 UI screenshot is claimed because the Qt interface is unchanged."),
    ])
# R112: enforce one absolute deadline across the Zabbix HTTPS transaction.
if (dest / "V31.8-R112-ZABBIX-TRANSACTION-DEADLINE.md").is_file():
    sections.extend([
        ("R112 - Zabbix absolute HTTPS deadline", "One monotonic deadline now spans HTTPS opening, TCP connect, TLS handshake, HTTP status/headers and body reading. A deadline watchdog shuts down the active socket; the body reader reuses the original deadline rather than receiving a new timeout budget."),
        ("R112 - TLS and privacy", "The default SSLContext still requires certificate validation and hostname checks; redirects remain rejected, response bytes remain capped at 1 MiB, and only the existing read-only problem.get and trigger.get methods are allowed. Errors contain no remote text or credential."),
        ("R112 - qualification boundary", "Seven new tests include a real loopback TLS server with an ephemeral runtime-generated trust anchor and slow-drip response headers. The temporary private key is deleted and is never archived. Platform getaddrinfo/DNS itself is not portably interruptible, and a real Zabbix instance, real token and least-privilege API role remain unverified. Zabbix stays PARTIAL; no R112 UI screenshot is claimed."),
    ])


# R113: bound caller waits for system DNS without creating unbounded threads.
if (dest / "V31.8-R113-ZABBIX-BOUNDED-DNS.md").is_file():
    sections.extend([
        ("R113 - bounded DNS transaction wait", "Zabbix HTTPS connection setup uses a fixed pool of two daemon DNS workers and a queue capped at two requests. The request thread shares the existing monotonic deadline and waits only until that deadline. Expired queued lookups are skipped before the OS resolver call begins. The first loopback CI run found that HTTPConnection assigns socket.create_connection on the instance and would shadow a same-named subclass method; R113 explicitly binds the instance callback to _deadline_create_connection and retains a regression test for that production path."),
        ("R113 - residual DNS limitation", "A native getaddrinfo call that has already started cannot be forcibly interrupted portably; it can remain in a worker. Fixed worker and queue limits prevent unbounded per-request thread creation, and saturated callers receive a neutral timeout."),
        ("R113 - tests and evidence", "Six unit tests cover expired deadline, valid getaddrinfo result shape, actual TCP loopback routing, bounded caller wait with a blocked resolver, expired queued lookup rejection and pool capacity limits. No external DNS, live Zabbix endpoint or real token is used. The cumulative workflow explicitly installs cryptography for the R112 ephemeral loopback TLS fixture, after its first run showed that the system Python lacked this test dependency. Windows native, OpenSSH and cumulative workflows must be verified at the exact commit before release. No R113 UI screenshot is claimed."),
    ])


# R101: packaged WinRM transport does not certify a real remote WinRM session.
if (dest / "V31.8-R101-WINRM-PACKAGED-RUNTIME.md").is_file():
    sections.extend([
        ("R101 - Windows WinRM packaged runtime", "The Windows build installs pywinrm 0.5.x and explicitly collects winrm and spnego modules in PyInstaller. The executable supports an offline dependency self-check without network or credential access."),
        ("R101 - binary acceptance gate", "The build executes the offline WinRM dependency self-check against both PyInstaller output and the independently extracted portable ZIP, with bounded timeouts. Missing modules block publication. Windows CI results must be checked at the exact commit."),
        ("R101 - unchanged security model", "WinRM remains disabled by default and uses only fixed read-only CIM queries over HTTPS with certificate validation. Credentials are resolved in memory after operator consent; errors are redacted."),
        ("R101 - unverified limits", "An importable WinRM dependency graph is not a successful connection to a real Windows Server. TLS trust, NTLM authentication, least-privilege remote access, physical Windows operator tests and real R101 screenshots remain unverified. The architecture SVG is conceptual, not a screenshot."),
    ])
# R114: Grafana dashboard discovery is opt-in and GET-only; synthetic tests do not attest live Grafana.
if (dest / "V31.8-R114-GRAFANA-READONLY.md").is_file():
    sections.extend([
        ("R114 - Grafana read-only dashboard discovery", "The Outils menu exposes an offline-by-default Grafana panel. It lists only dashboard search results returned by GET /api/search with type=dash-db, projects UID/title/folder/tags, bounds the response to 1 MiB and 100 dashboards, rejects redirects, and rebuilds links from a validated HTTPS origin and UID. The browser opens only after an explicit operator action."),
        ("R114 - consent and secret handling", "The origin and search query must remain unchanged after consent; edits, cancel, hide or close clear displayed data and revoke consent. OPENREMOTE_GRAFANA_TOKEN is read only after consent and a valid HTTPS origin. The bearer is sent in the Authorization header, not in URLs, logs, persistence or exports. Errors are normalized."),
        ("R114 - tests and evidence", "Eighteen backend contract tests and nine Qt tests use fake transports/clients and a test token; they make no request to an external Grafana service. Windows CI generates a real offscreen PNG of the Grafana panel with consent disabled; that image is evidence of the offline UI state only, not of a live dashboard."),
        ("R114 - limitations and release boundary", "A real Grafana instance, real token and effective least-privilege role remain unverified. Pagination, Grafana reverse-proxy subpaths, physical operator acceptance and signed binaries are not qualified. The transport uses strict urllib TLS trust, rejects redirects and enforces response limits; it cannot guarantee immediate interruption of an already blocked native read. Grafana is PARTIAL, not IMPLEMENTED."),
    ])

doc = Document()
doc.add_heading("OpenRemote Studio V" + version, 0)
doc.add_paragraph("Windows release quality dossier | Abdelhalim Talbi")
for heading, body in sections:
    doc.add_heading(heading, level=1)
    doc.add_paragraph(body)
doc.save(dest / ("Dossier-V" + version + ".docx"))
styles = getSampleStyleSheet()
story = [Paragraph(escape("OpenRemote Studio V" + version), styles["Title"]), Spacer(1, 14)]
for heading, body in sections:
    story.extend([Paragraph(escape(heading), styles["Heading2"]), Paragraph(escape(body), styles["BodyText"]), Spacer(1, 12)])
SimpleDocTemplate(str(dest / ("Dossier-V" + version + ".pdf"))).build(story)
print("Generated versioned DOCX/PDF release dossier for V" + version)
